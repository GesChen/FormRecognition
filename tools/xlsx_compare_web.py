#!/usr/bin/env python3
"""
Temporary web tool for XLSX truth-vs-test comparison.

Run:
  python3 tools/xlsx_compare_web.py [--port 5004]
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import threading
import uuid
from pathlib import Path

from flask import Flask, jsonify, request, send_file
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from werkzeug.utils import secure_filename

TOOLS_DIR = Path(__file__).resolve().parent
ROOT = TOOLS_DIR.parent

# Reuse temporary comparer logic under testing/.
if str(ROOT / "testing") not in sys.path:
    sys.path.insert(0, str(ROOT / "testing"))
import tmp_xlsx_comparer as comparer  # type: ignore


UPLOAD_DIR = ROOT / "testing" / "output" / "web_tools" / "xlsx_compare_uploads"
ALLOWED_SUFFIXES = {".xlsx", ".xlsm"}
DEFAULT_TRUTH_DIR = ROOT / "data" / "2025 NPS FLE Session Delivery"
DEFAULT_TEST_DIR = ROOT / "output" / "xlsx"

_uploads_lock = threading.Lock()
_uploads: dict[str, Path] = {}
SERVER_INSTANCE_ID = uuid.uuid4().hex


def _is_allowed_xlsx(name: str) -> bool:
    return Path(name or "").suffix.lower() in ALLOWED_SUFFIXES


def _sheets_for_xlsx(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _resolve_upload(upload_id: str) -> Path | None:
    with _uploads_lock:
        p = _uploads.get(upload_id)
    if not p or not p.is_file():
        return None
    return p


def _resolve_default_path(path_text: str, default_dir: Path) -> Path | None:
    if not path_text:
        return None
    try:
        path = Path(path_text).resolve()
    except OSError:
        return None
    if not path.is_file() or path.suffix.lower() not in ALLOWED_SUFFIXES:
        return None
    return path


def _default_dir_for_role(role: str) -> Path | None:
    if role == "truth":
        return DEFAULT_TRUTH_DIR
    if role == "test":
        return DEFAULT_TEST_DIR
    return None


def _workbook_upload_payload(path: Path) -> dict[str, object]:
    sheets = _sheets_for_xlsx(path)
    upload_id = str(uuid.uuid4())
    with _uploads_lock:
        _uploads[upload_id] = path
    return {
        "ok": True,
        "upload_id": upload_id,
        "filename": path.name,
        "saved_path": str(path),
        "sheets": sheets,
    }


def _pick_xlsx_file(default_dir: Path, title: str) -> Path | None:
    if not default_dir.is_dir():
        raise RuntimeError(f"Default directory does not exist: {default_dir}")

    kdialog = shutil.which("kdialog")
    if kdialog:
        proc = subprocess.run(
            [
                kdialog,
                "--title",
                title,
                "--getopenfilename",
                str(default_dir),
                "*.xlsx *.xlsm|Excel workbooks (*.xlsx *.xlsm)",
            ],
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        if proc.returncode == 1:
            return None
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr.strip() or "File picker failed")
        picked = proc.stdout.strip()
        return Path(picked) if picked else None

    zenity = shutil.which("zenity")
    if zenity:
        proc = subprocess.run(
            [
                zenity,
                "--file-selection",
                f"--title={title}",
                f"--filename={default_dir.as_posix()}/",
                "--file-filter=Excel workbooks | *.xlsx *.xlsm",
            ],
            check=False,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        if proc.returncode == 1:
            return None
        if proc.returncode != 0:
            raise RuntimeError(proc.stderr.strip() or "File picker failed")
        picked = proc.stdout.strip()
        return Path(picked) if picked else None

    try:
        import tkinter as tk
        from tkinter import filedialog
    except Exception as exc:
        raise RuntimeError("No supported native file picker found (kdialog, zenity, or tkinter)") from exc

    root = tk.Tk()
    root.withdraw()
    try:
        picked = filedialog.askopenfilename(
            title=title,
            initialdir=str(default_dir),
            filetypes=[("Excel workbooks", "*.xlsx *.xlsm")],
        )
    finally:
        root.destroy()
    return Path(picked) if picked else None


def _browse_xlsx_files(role: str, rel_dir: str = "") -> tuple[list[str], list[str]]:
    """List subdirs and .xlsx/.xlsm files for a given role (truth/test)."""
    default_dir = _default_dir_for_role(role)
    if default_dir is None or not default_dir.exists():
        return [], []

    rel_dir = (rel_dir or "").strip().rstrip("/")
    if rel_dir and (".." in rel_dir or rel_dir.startswith("/")):
        return [], []

    target = (default_dir / rel_dir).resolve() if rel_dir else default_dir
    try:
        target.relative_to(default_dir.resolve())
    except ValueError:
        return [], []

    if not target.is_dir():
        return [], []

    dirs = []
    files = []
    try:
        for p in target.iterdir():
            try:
                if p.is_dir():
                    rel = p.relative_to(default_dir)
                    dirs.append(str(rel).replace("\\", "/"))
                elif p.is_file() and p.suffix.lower() in ALLOWED_SUFFIXES:
                    rel = p.relative_to(default_dir)
                    files.append(str(rel).replace("\\", "/"))
            except (ValueError, OSError):
                continue
    except (OSError, PermissionError):
        pass

    return sorted(dirs), sorted(files)


def _last_nonempty_id_row(ws, start_row: int, id_col_idx: int) -> int:
    last = start_row
    max_row = int(ws.max_row or start_row)
    for r in range(start_row, max_row + 1):
        v = ws.cell(r, id_col_idx).value
        if v is not None and str(v).strip() != "":
            last = r
    return last


def create_app() -> Flask:
    app = Flask(
        __name__,
        static_folder=TOOLS_DIR / "static",
        template_folder=TOOLS_DIR / "templates",
    )

    @app.route("/")
    def index():
        return send_file(TOOLS_DIR / "templates" / "xlsx_compare.html")

    @app.route("/api/instance")
    def api_instance():
        return jsonify({"instance_id": SERVER_INSTANCE_ID})

    @app.route("/api/browse")
    def api_browse():
        role = request.args.get("role", "").strip().lower()
        rel_dir = request.args.get("dir", "").strip()
        if role not in ("truth", "test"):
            return jsonify({"error": "role must be 'truth' or 'test'"}), 400
        dirs, files = _browse_xlsx_files(role, rel_dir)
        return jsonify({"dirs": dirs, "files": files})

    @app.route("/api/upload-xlsx", methods=["POST"])
    def api_upload_xlsx():
        up = request.files.get("file")
        if not up or not up.filename:
            return jsonify({"error": "Missing file"}), 400
        if not _is_allowed_xlsx(up.filename):
            return jsonify({"error": "Only .xlsx/.xlsm are supported"}), 400

        role = (request.form.get("role") or "file").strip().lower()
        safe = secure_filename(up.filename) or "workbook.xlsx"
        suffix = Path(safe).suffix.lower()
        if suffix not in ALLOWED_SUFFIXES:
            suffix = ".xlsx"
        stem = Path(safe).stem or "workbook"
        upload_id = str(uuid.uuid4())
        out_name = f"{role}_{upload_id}_{stem}{suffix}"

        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        out_path = (UPLOAD_DIR / out_name).resolve()
        up.save(str(out_path))

        try:
            sheets = _sheets_for_xlsx(out_path)
        except Exception as exc:
            try:
                out_path.unlink(missing_ok=True)
            except OSError:
                pass
            return jsonify({"error": f"Could not read workbook: {exc}"}), 400

        with _uploads_lock:
            _uploads[upload_id] = out_path

        try:
            rel = out_path.relative_to(ROOT.resolve()).as_posix()
        except ValueError:
            rel = str(out_path)

        return jsonify(
            {
                "ok": True,
                "upload_id": upload_id,
                "filename": up.filename,
                "saved_path": rel,
                "sheets": sheets,
            }
        )

    @app.route("/api/load-workbook", methods=["POST"])
    def api_load_workbook():
        body = request.get_json(silent=True) or {}
        role = str(body.get("role", "")).strip().lower()
        file_path = str(body.get("file_path", "")).strip()
        if role not in ("truth", "test"):
            return jsonify({"error": "role must be 'truth' or 'test'"}), 400
        if not file_path:
            return jsonify({"error": "file_path is required"}), 400

        default_dir = _default_dir_for_role(role)
        if default_dir is None:
            return jsonify({"error": "Invalid role"}), 400

        workbook_path = (default_dir / file_path).resolve()
        try:
            workbook_path.relative_to(default_dir.resolve())
        except ValueError:
            return jsonify({"error": "Path escapes default directory"}), 400

        if not workbook_path.is_file() or workbook_path.suffix.lower() not in ALLOWED_SUFFIXES:
            return jsonify({"error": f"File not found or not an Excel file: {file_path}"}), 400

        try:
            payload = _workbook_upload_payload(workbook_path)
        except Exception as exc:
            return jsonify({"error": f"Could not read workbook: {exc}"}), 400
        return jsonify(payload)

    @app.route("/api/pick-default-workbook", methods=["POST"])
    def api_pick_default_workbook():
        body = request.get_json(silent=True) or {}
        role = str(body.get("role", "")).strip().lower()
        default_dir = _default_dir_for_role(role)
        if default_dir is None:
            return jsonify({"error": "role must be 'truth' or 'test'"}), 400

        try:
            picked = _pick_xlsx_file(default_dir, f"Select {role} workbook")
        except Exception as exc:
            return jsonify({"error": str(exc)}), 500
        if picked is None:
            return jsonify({"ok": False, "cancelled": True})

        workbook_path = _resolve_default_path(str(picked), default_dir)
        if not workbook_path:
            return jsonify({"error": f"Selected {role} workbook must be an existing .xlsx/.xlsm file"}), 400

        try:
            payload = _workbook_upload_payload(workbook_path)
        except Exception as exc:
            return jsonify({"error": f"Could not read workbook: {exc}"}), 400
        return jsonify(payload)

    @app.route("/api/compare", methods=["POST"])
    def api_compare():
        body = request.get_json(silent=True) or {}
        truth_upload_id = str(body.get("truth_upload_id", "")).strip()
        test_upload_id = str(body.get("test_upload_id", "")).strip()
        truth_sheet = str(body.get("truth_sheet", "")).strip()
        test_sheet = str(body.get("test_sheet", "")).strip()
        id_match_any_truth_sheet = bool(body.get("id_match_any_truth_sheet", False))

        if not truth_upload_id or not test_upload_id:
            return jsonify({"error": "Both truth and test uploads are required"}), 400
        if not test_sheet:
            return jsonify({"error": "test_sheet is required"}), 400
        if (not id_match_any_truth_sheet) and (not truth_sheet):
            return jsonify({"error": "truth_sheet is required unless ID-match-any-truth-sheet is enabled"}), 400

        truth_file = _resolve_upload(truth_upload_id)
        test_file = _resolve_upload(test_upload_id)
        if not truth_file or not test_file:
            return jsonify({"error": "One or both uploaded files are missing; re-upload and retry"}), 400

        try:
            start_row = int(body.get("start_row", 4))
        except (TypeError, ValueError):
            return jsonify({"error": "start_row must be an integer"}), 400
        end_row_raw = body.get("end_row")
        if end_row_raw in (None, ""):
            end_row = None
        else:
            try:
                end_row = int(end_row_raw)
            except (TypeError, ValueError):
                return jsonify({"error": "end_row must be an integer when provided"}), 400

        start_col_label = str(body.get("start_col", "A")).strip().upper() or "A"
        end_col_label = str(body.get("end_col", "")).strip().upper()
        id_col_label = str(body.get("id_column", "A")).strip().upper() or "A"
        id_regex_text = str(body.get("id_regex", r"^ID:[A-Z0-9]{8}$")).strip() or r"^ID:[A-Z0-9]{8}$"

        try:
            start_col = column_index_from_string(start_col_label)
            id_col = column_index_from_string(id_col_label)
        except ValueError as exc:
            return jsonify({"error": f"Invalid column label: {exc}"}), 400
        try:
            id_regex = re.compile(id_regex_text, flags=re.IGNORECASE)
        except re.error as exc:
            return jsonify({"error": f"Invalid id_regex: {exc}"}), 400

        wb_truth = None
        wb_test = None
        try:
            wb_truth = load_workbook(truth_file, data_only=True)
            wb_test = load_workbook(test_file, data_only=True)
            if (not id_match_any_truth_sheet) and truth_sheet not in wb_truth.sheetnames:
                return jsonify({"error": f"Truth sheet not found: {truth_sheet!r}"}), 400
            if test_sheet not in wb_test.sheetnames:
                return jsonify({"error": f"Test sheet not found: {test_sheet!r}"}), 400
            ws_truth = wb_truth[truth_sheet] if not id_match_any_truth_sheet else None
            ws_test = wb_test[test_sheet]
            if end_col_label:
                end_col = column_index_from_string(end_col_label)
            else:
                if ws_truth is None:
                    end_col = int(ws_test.max_column or 1)
                else:
                    end_col = max(int(ws_truth.max_column or 1), int(ws_test.max_column or 1))
            if end_row is None:
                # Compare scope follows rows present in the test workbook.
                end_row = _last_nonempty_id_row(ws_test, start_row, id_col)
        except ValueError as exc:
            return jsonify({"error": f"Invalid compare range: {exc}"}), 400
        except Exception as exc:
            return jsonify({"error": f"Could not prepare workbooks: {exc}"}), 400
        finally:
            if wb_truth is not None:
                wb_truth.close()
            if wb_test is not None:
                wb_test.close()

        if end_col < start_col:
            return jsonify({"error": "end_col must be >= start_col"}), 400
        if end_row < start_row:
            return jsonify({"error": "end_row must be >= start_row"}), 400

        try:
            result = comparer.compare(
                truth_file=truth_file,
                truth_sheet=(truth_sheet or None),
                test_file=test_file,
                test_sheet=test_sheet,
                compare_range=comparer.CompareRange(
                    start_row=start_row,
                    end_row=end_row,
                    start_col=start_col,
                    end_col=end_col,
                ),
                id_col_idx=id_col,
                id_regex=id_regex,
                id_match_any_truth_sheet=id_match_any_truth_sheet,
            )
        except Exception as exc:
            return jsonify({"error": f"Compare failed: {exc}"}), 500

        return jsonify({"ok": True, "result": result})

    return app


def run(port: int = 5004, debug: bool = False) -> None:
    app = create_app()
    print(f"XLSX compare tool: http://127.0.0.1:{port}/")
    app.run(host="0.0.0.0", port=port, debug=debug)


if __name__ == "__main__":
    p = argparse.ArgumentParser(description="Temporary XLSX truth-vs-test comparer")
    p.add_argument("--port", type=int, default=5004, help="Port (default: 5004)")
    p.add_argument("--debug", action="store_true", help="Flask debug mode")
    args = p.parse_args()
    run(port=args.port, debug=args.debug)
