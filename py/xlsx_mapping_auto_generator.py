"""Auto-generate XLSX mapping JSON files from templates + ROI context."""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Mapping

from llm_client import generate
from openai_client import call_json, resolve_model
from form_sheet_map import form_sheet_name, workbook_template_path
from roi_auto_detector import _collect_document_structure, _extract_json_object

try:
    from config import (
        PATHS,
        PROJECT_ROOT,
        XLSX_DATA_ENTRY,
        XLSX_MAPPING_AUTO_GENERATE,
    )
except Exception:  # pragma: no cover
    ROOT = Path(__file__).resolve().parent.parent
    PROJECT_ROOT = ROOT
    PATHS = {
        "data": ROOT / "data",
        "templates_root": ROOT / "data" / "templates",
        "roi_schemas_root": ROOT / "data" / "roi_schemas",
    }
    XLSX_DATA_ENTRY = {}
    XLSX_MAPPING_AUTO_GENERATE = {
        "use_openai": True,
        "openai_model": "gpt-5.5",
        "local_model": "qwen3.5:9b",
    }

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    load_workbook = None  # type: ignore[assignment]
    get_column_letter = None  # type: ignore[assignment]
    _OPENPYXL_ERR = exc
else:
    _OPENPYXL_ERR = None

FORM_NAME_RE = re.compile(r"^[A-Za-z0-9_-]+$")
RESERVED_SOURCES = {"id", "form_type", "page_odd", "page_even"}


class XlsxMappingAutoGenerationError(RuntimeError):
    """Raised with phase context for otherwise opaque auto-generation failures."""


def _phase_error(phase: str, exc: Exception) -> XlsxMappingAutoGenerationError:
    msg = str(exc) or type(exc).__name__
    return XlsxMappingAutoGenerationError(f"{phase}: {type(exc).__name__}: {msg}")


def _require_openpyxl() -> None:
    if _OPENPYXL_ERR is not None:
        raise ImportError("openpyxl is required for XLSX mapping auto-generation.") from _OPENPYXL_ERR


def _resolve_collection_root(path_value: Path | str, expected_leaf: str) -> Path:
    p = Path(path_value).resolve()
    if p.name == expected_leaf:
        return p
    if p.parent.name == expected_leaf:
        return p.parent
    return p


DATA_ROOT = Path(PATHS.get("data", Path(PROJECT_ROOT) / "data")).resolve()
TEMPLATES_ROOT = _resolve_collection_root(PATHS.get("templates_root", DATA_ROOT / "templates"), "templates")
ROI_SCHEMAS_ROOT = _resolve_collection_root(PATHS.get("roi_schemas_root", DATA_ROOT / "roi_schemas"), "roi_schemas")


def _safe_name(value: str, label: str) -> str:
    s = str(value or "").strip()
    if not s or not FORM_NAME_RE.match(s):
        raise ValueError(f"Invalid {label}: {value!r}")
    return s


def _project_rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(Path(PROJECT_ROOT).resolve())).replace("\\", "/")
    except Exception:
        return str(path.resolve())


def resolve_local_xlsx_path(path_value: str | Path | None = None, *, release: str | None = None) -> Path:
    """Resolve a backend-readable XLSX path, defaulting to the release template."""
    raw = str(path_value or "").strip()
    if not raw:
        return workbook_template_path(release)
    # Explicit paths are retained for tests and non-UI scripts; web tools use
    # release-based template resolution instead of manual XLSX selection.
    if not raw:
        raise FileNotFoundError("No template workbook path provided.")
    p = Path(raw).expanduser()
    if not p.is_absolute():
        p = Path(PROJECT_ROOT).resolve() / p
    p = p.resolve()
    if not p.is_file():
        raise FileNotFoundError(f"Target template workbook not found: {p}")
    if p.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ValueError(f"Target template path is not an XLSX workbook: {p}")
    return p


def _configured_sheet_name(form: str, release: str | None = None) -> str | None:
    return form_sheet_name(form, release)


def _resolve_sheet_name(workbook_path: Path, form: str, sheet_hint: str | None = None, release: str | None = None) -> str:
    _require_openpyxl()
    wb = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheetnames = list(wb.sheetnames)
    finally:
        wb.close()
    for candidate in (sheet_hint, _configured_sheet_name(form, release)):
        if candidate and candidate in sheetnames:
            return str(candidate)
    normalized_form = form.lower().replace("_", "")
    grade = "HS" if normalized_form.startswith("h") else normalized_form[:1]
    timing = "post" if "post" in normalized_form else "pre"
    for name in sheetnames:
        low = name.lower()
        if timing in low and (grade.lower() in low or (grade == "HS" and "hs" in low)):
            return name
    if sheetnames:
        return sheetnames[0]
    raise ValueError(f"Workbook has no sheets: {workbook_path}")


def _cell_value(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).split())


def extract_sheet_structure(
    workbook_path: Path,
    *,
    form: str,
    release: str | None = None,
    sheet_hint: str | None = None,
    start_row_hint: int | None = None,
) -> dict[str, Any]:
    """Return a compact, prompt-friendly structure for one workbook sheet."""
    _require_openpyxl()
    sheet_name = _resolve_sheet_name(workbook_path, form, sheet_hint=sheet_hint, release=release)
    start_row = int(start_row_hint or XLSX_MAPPING_AUTO_GENERATE.get("default_start_row", 4) or 4)
    max_header_row = max(start_row, int(XLSX_MAPPING_AUTO_GENERATE.get("xlsx_context_rows", 6) or 6))

    wb = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        ws = wb[sheet_name]
        max_col = int(ws.max_column or 1)
        max_row = int(ws.max_row or 1)
        columns: list[dict[str, Any]] = []
        for col_idx in range(1, max_col + 1):
            letter = get_column_letter(col_idx)
            values: list[dict[str, Any]] = []
            for row_idx in range(1, min(max_header_row, max_row) + 1):
                text = _cell_value(ws.cell(row=row_idx, column=col_idx).value)
                if text:
                    values.append({"row": row_idx, "value": text})
            columns.append(
                {
                    "column": letter,
                    "header_values": values,
                    "start_row_value": _cell_value(ws.cell(row=start_row, column=col_idx).value),
                    "width": ws.column_dimensions[letter].width,
                }
            )
        merged_ranges: list[str] = []
        for merged in ws.merged_cells.ranges:
            if merged.min_row <= max_header_row:
                merged_ranges.append(str(merged))
    finally:
        wb.close()

    return {
        "workbook_path": _project_rel(workbook_path),
        "sheet_name": sheet_name,
        "max_row": max_row,
        "max_column": max_col,
        "suggested_start_row": start_row,
        "columns": columns,
        "merged_ranges_in_header_area": merged_ranges,
    }


def _load_json_file(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def _compact_roi_schema(path: Path, side: str) -> dict[str, Any]:
    data = _load_json_file(path)
    rois: list[dict[str, Any]] = []
    for roi in data.get("rois") or []:
        if not isinstance(roi, dict):
            continue
        name = str(roi.get("name", "")).strip()
        if not name:
            continue
        rois.append(
            {
                "name": name,
                "x": int(float(roi.get("x", 0) or 0)),
                "y": int(float(roi.get("y", 0) or 0)),
                "w": int(float(roi.get("w", 0) or 0)),
                "h": int(float(roi.get("h", 0) or 0)),
            }
        )
    return {
        "side": side,
        "schema_path": _project_rel(path),
        "image_width": data.get("image_width"),
        "image_height": data.get("image_height"),
        "rois": rois,
    }


def load_target_roi_pair(release: str, form: str) -> dict[str, Any]:
    release = _safe_name(release, "release")
    form = _safe_name(form, "form")
    sides: dict[str, Any] = {}
    all_sources = set(RESERVED_SOURCES)
    for side in ("a", "b"):
        path = ROI_SCHEMAS_ROOT / release / f"{form}_{side}.json"
        if not path.is_file():
            raise FileNotFoundError(f"ROI schema not found for {form}_{side}: {path}")
        compact = _compact_roi_schema(path, side)
        sides[side] = compact
        all_sources.update(str(r.get("name", "")).strip() for r in compact.get("rois", []) if r.get("name"))
    return {
        "release": release,
        "form": form,
        "reserved_sources": sorted(RESERVED_SOURCES),
        "all_sources": sorted(all_sources, key=_source_sort_key),
        "sides": sides,
    }


def _source_sort_key(value: str) -> tuple[int, int, str]:
    s = str(value)
    if s.isdigit():
        return (0, int(s), s)
    m = re.match(r"^(\d+)([A-Za-z]+)$", s)
    if m:
        return (1, int(m.group(1)), m.group(2).lower())
    if s in RESERVED_SOURCES:
        return (-1, 0, s)
    return (2, 0, s.lower())


def collect_target_paddle_pair(release: str, form: str) -> dict[str, Any]:
    release = _safe_name(release, "release")
    form = _safe_name(form, "form")
    sides: dict[str, Any] = {}
    for side in ("a", "b"):
        image_path = TEMPLATES_ROOT / release / f"{form}_{side}.png"
        if not image_path.is_file():
            raise FileNotFoundError(f"Template image not found for {form}_{side}: {image_path}")
        try:
            document_structure = _collect_document_structure(image_path)
        except Exception as exc:
            raise _phase_error(
                f"Paddle/template OCR for {form}_{side} ({_project_rel(image_path)})",
                exc,
            ) from exc
        sides[side] = {
            "template_image_path": _project_rel(image_path),
            "document_structure_from_paddleocr": document_structure,
        }
    return {"release": release, "form": form, "sides": sides}


def _example_dir() -> Path:
    raw = XLSX_MAPPING_AUTO_GENERATE.get("example_dir", DATA_ROOT / "xlsx" / "mappings" / "example" / "6pre")
    p = Path(raw)
    if not p.is_absolute():
        p = Path(PROJECT_ROOT).resolve() / p
    return p.resolve()


def load_example_package() -> dict[str, Any]:
    cfg = XLSX_MAPPING_AUTO_GENERATE if isinstance(XLSX_MAPPING_AUTO_GENERATE, dict) else {}
    example_form = str(cfg.get("example_form", "6pre") or "6pre")
    root = _example_dir()
    mapping_path = Path(cfg.get("example_mapping_path") or root / "mapping.json")
    template_path = Path(cfg.get("example_template_workbook") or root / "template.xlsx")
    paddle_a_path = Path(cfg.get("example_paddle_a_path") or root / "paddle_a.json")
    paddle_b_path = Path(cfg.get("example_paddle_b_path") or root / "paddle_b.json")
    paths = [mapping_path, template_path, paddle_a_path, paddle_b_path]
    resolved: list[Path] = []
    for path in paths:
        p = path if path.is_absolute() else Path(PROJECT_ROOT).resolve() / path
        if not p.exists():
            raise FileNotFoundError(f"Example asset not found: {p}")
        resolved.append(p.resolve())
    mapping = _load_json_file(resolved[0])
    sheet = str(mapping.get("sheet", "") or "")
    start_row = int(mapping.get("start_row", 4) or 4)
    return {
        "form": str(mapping.get("form_type") or example_form),
        "mapping_path": _project_rel(resolved[0]),
        "template_workbook_path": _project_rel(resolved[1]),
        "paddle_a_path": _project_rel(resolved[2]),
        "paddle_b_path": _project_rel(resolved[3]),
        "mapping": mapping,
        "paddle_pair": {
            "a": _load_json_file(resolved[2]),
            "b": _load_json_file(resolved[3]),
        },
        "xlsx_sheet_structure": extract_sheet_structure(
            resolved[1],
            form=str(mapping.get("form_type") or example_form),
            sheet_hint=sheet,
            start_row_hint=start_row,
        ),
    }


def _generation_schema() -> dict[str, Any]:
    value_schema = {"type": ["string", "number", "boolean", "null"]}
    entry_schema = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "source": {"type": ["string", "null"]},
            "type": {"type": "string", "enum": ["direct", "lookup", "multi_column", "static"]},
            "column": {"type": ["string", "null"]},
            "transform": {"type": ["string", "null"], "enum": ["upper", "lower", "number", "date", "", None]},
            "map_entries": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {"key": {"type": "string"}, "value": value_schema},
                    "required": ["key", "value"],
                },
            },
            "default": value_schema,
            "choices": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {"key": {"type": "string"}, "column": {"type": "string"}},
                    "required": ["key", "column"],
                },
            },
            "mark": value_schema,
            "no_answer": {"type": ["string", "null"]},
            "blank": value_schema,
            "value": value_schema,
            "rationale": {"type": "string"},
        },
        "required": [
            "source",
            "type",
            "column",
            "transform",
            "map_entries",
            "default",
            "choices",
            "mark",
            "no_answer",
            "blank",
            "value",
            "rationale",
        ],
    }
    return {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "form_type": {"type": "string"},
            "sheet": {"type": "string"},
            "start_row": {"type": "integer", "minimum": 1},
            "mappings": {"type": "array", "minItems": 1, "items": entry_schema},
            "confidence_notes": {"type": "array", "items": {"type": "string"}},
        },
        "required": ["form_type", "sheet", "start_row", "mappings", "confidence_notes"],
    }


def _instructions() -> str:
    return (
        "You generate complete XLSX mapping JSON for the EVMS data-entry pipeline.\n"
        "Return ONE JSON object only that matches the provided schema exactly. No markdown, prose, comments, or extra keys.\n"
        "The generated mapping must be production-ready and must be correct.\n\n"
        "Mapping semantics:\n"
        "- direct: write one source value to one Excel column. Use transform=upper for MCQ answer letters, number for age/number fields, date for normalized dates when appropriate.\n"
        "- lookup: map one source answer letter/value through map_entries into one Excel column.\n"
        "- multi_column: one source answer selects exactly one target column from choices; mark is usually Yes; no_answer is the column for explicit no-answer when present.\n"
        "- static: write a constant value to one target column.\n\n"
        "Source rules:\n"
        "- Use only target_roi_schemas.all_sources plus reserved sources id, form_type, page_odd, page_even.\n"
        "- MCQ question sources are the base question names (for example source '22'), not choice ROI names ('22a', '22b'), unless the workbook truly needs per-choice filled flags.\n"
        "- Do not invent a source that is absent from target ROI schemas.\n\n"
        "Workbook rules:\n"
        "- Use target_xlsx_sheet_structure as the source of truth for sheet name, start row, and target columns.\n"
        "- Do not output columns outside target_xlsx_sheet_structure.columns.\n"
        "- Cover every target workbook data column that should be filled by recognition output for this form; do not create placeholder mappings.\n"
        "- Preserve workbook-specific answer labels exactly where headers imply specific text (TRUE/FALSE, Boy/Girl, grade names, Yes marks, etc.).\n\n"
        "Anti-overfitting rules for the example package:\n"
        "- The example package is a convention guide, not an answer key for the target.\n"
        "- Do not copy example columns, question counts, source names, sheet names, or lookup values unless the target workbook structure and target ROI/paddle evidence independently support them.\n"
        "- If the target has the same-looking workbook as the example, still verify each target column against target headers before using it.\n"
        "- Prefer target workbook headers over example mapping whenever they conflict.\n"
        "- Prefer target ROI names and target Paddle text over example Paddle whenever they conflict.\n\n"
        "Correctness checks before final output:\n"
        "- Every mapping row has a valid type.\n"
        "- direct/lookup/static rows have a single valid column.\n"
        "- lookup rows include non-empty map_entries.\n"
        "- multi_column rows include non-empty choices with valid columns and a non-empty mark.\n"
        "- All source names exist in target_roi_schemas.all_sources unless row type is static.\n"
        "- The form_type equals target.form and sheet equals the resolved target sheet.\n"
        "- confidence_notes should list any uncertainty; do not hide uncertainty by inventing.\n"
    )


def _normalize_generated_mapping(raw_obj: dict[str, Any], *, form: str, fallback_sheet: str, fallback_start_row: int) -> dict[str, Any]:
    if not isinstance(raw_obj, dict):
        raise ValueError("Generated result must be a JSON object.")
    rows = raw_obj.get("mappings")
    if not isinstance(rows, list) or not rows:
        raise ValueError("Generated mapping has no rows.")
    out_rows: list[dict[str, Any]] = []
    for idx, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            raise ValueError(f"Generated row {idx} is not an object.")
        mtype = str(row.get("type", "")).strip().lower()
        if mtype not in {"direct", "lookup", "multi_column", "static"}:
            raise ValueError(f"Generated row {idx} has invalid type {mtype!r}.")
        out: dict[str, Any] = {"type": mtype}
        source = row.get("source")
        if mtype != "static":
            out["source"] = str(source or "").strip()
        column = str(row.get("column") or "").strip().upper()
        if mtype in {"direct", "lookup", "static"}:
            out["column"] = column
        transform = str(row.get("transform") or "").strip().lower()
        if mtype == "direct" and transform:
            out["transform"] = transform
        if mtype == "lookup":
            entries = row.get("map_entries") or []
            mp: dict[str, Any] = {}
            for item in entries:
                if not isinstance(item, dict):
                    continue
                key = str(item.get("key", "")).strip()
                if key:
                    mp[key] = item.get("value")
            out["map"] = mp
            default = row.get("default")
            if default is not None:
                out["default"] = default
        if mtype == "multi_column":
            choices = row.get("choices") or []
            choice_map: dict[str, str] = {}
            for item in choices:
                if not isinstance(item, dict):
                    continue
                key = str(item.get("key", "")).strip()
                col = str(item.get("column", "")).strip().upper()
                if key and col:
                    choice_map[key] = col
            out["choices"] = choice_map
            mark = row.get("mark")
            out["mark"] = "Yes" if mark in (None, "") else mark
            no_answer = str(row.get("no_answer") or "").strip().upper()
            if no_answer:
                out["no_answer"] = no_answer
            blank = row.get("blank")
            if blank is not None:
                out["blank"] = blank
        if mtype == "static":
            out["value"] = row.get("value")
        out_rows.append(out)

    return {
        "form_type": form,
        "sheet": str(raw_obj.get("sheet") or fallback_sheet),
        "start_row": int(raw_obj.get("start_row") or fallback_start_row),
        "mappings": out_rows,
    }


def _mapping_destinations(row: dict[str, Any]) -> list[str]:
    mtype = str(row.get("type", "")).strip().lower()
    cols: list[str] = []
    if mtype in {"direct", "lookup", "static"}:
        col = str(row.get("column", "") or "").strip().upper()
        if col:
            cols.append(col)
    elif mtype == "multi_column":
        choices = row.get("choices") if isinstance(row.get("choices"), dict) else {}
        for col in choices.values():
            c = str(col or "").strip().upper()
            if c:
                cols.append(c)
        no_answer = str(row.get("no_answer", "") or "").strip().upper()
        if no_answer:
            cols.append(no_answer)
    return cols


def _workbook_mapping_issues(mapping: dict[str, Any], target_xlsx: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    allowed_cols = {
        str(col.get("column", "")).strip().upper()
        for col in target_xlsx.get("columns", [])
        if isinstance(col, dict) and col.get("column")
    }
    expected_sheet = str(target_xlsx.get("sheet_name", "") or "")
    if str(mapping.get("sheet", "") or "") != expected_sheet:
        issues.append(
            {
                "severity": "error",
                "row": None,
                "message": f"Generated sheet {mapping.get('sheet')!r} does not match target sheet {expected_sheet!r}.",
            }
        )
    expected_start = int(target_xlsx.get("suggested_start_row", 0) or 0)
    try:
        got_start = int(mapping.get("start_row", 0) or 0)
    except Exception:
        got_start = 0
    if expected_start and got_start != expected_start:
        issues.append(
            {
                "severity": "error",
                "row": None,
                "message": f"Generated start_row {got_start!r} does not match target start row {expected_start}.",
            }
        )
    for idx, row in enumerate(mapping.get("mappings") or [], start=1):
        if not isinstance(row, dict):
            continue
        for col in _mapping_destinations(row):
            if col not in allowed_cols:
                issues.append(
                    {
                        "severity": "error",
                        "row": idx,
                        "message": f"Column {col!r} is outside target workbook sheet columns.",
                    }
                )
    return issues


def _call_local_json(payload: dict[str, Any], schema: dict[str, Any], instructions: str) -> dict[str, Any]:
    cfg = XLSX_MAPPING_AUTO_GENERATE if isinstance(XLSX_MAPPING_AUTO_GENERATE, dict) else {}
    model = str(cfg.get("local_model", "qwen3.5:9b") or "qwen3.5:9b")
    reasoning = bool(cfg.get("local_reasoning", False))
    prompt = (
        f"{instructions}\n\n"
        "Strict output requirements:\n"
        "- Return exactly one JSON object.\n"
        "- No markdown, no code fences, no explanation.\n"
        "- JSON must match this schema:\n"
        f"{json.dumps(schema, ensure_ascii=False)}\n\n"
        "Input payload:\n"
        f"{json.dumps(payload, ensure_ascii=False)}"
    )
    extra_params: dict[str, Any] = {"options": {"temperature": 0}, "format": schema}
    if reasoning:
        extra_params["think"] = True
    out = generate(
        prompt,
        model=model,
        timeout=int(cfg.get("local_timeout_sec", 600) or 600),
        extra_params=extra_params,
    )
    text = str(out.get("text", "") or "")
    if not text.strip():
        raise RuntimeError("Local model returned empty output for XLSX mapping generation.")
    return _extract_json_object(text)


def generate_xlsx_mapping(
    *,
    release: str,
    form: str,
    template_workbook_path: str | Path | None = None,
    sheet_hint: str | None = None,
    start_row_hint: int | None = None,
    validation_feedback: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    """Generate one complete mapping and return mapping + context metadata."""
    release = _safe_name(release, "release")
    form = _safe_name(form, "form")
    try:
        workbook_path = resolve_local_xlsx_path(template_workbook_path, release=release)
    except (FileNotFoundError, ValueError):
        raise
    except Exception as exc:
        raise _phase_error("Resolve template workbook path", exc) from exc
    try:
        target_xlsx = extract_sheet_structure(
            workbook_path,
            form=form,
            release=release,
            sheet_hint=sheet_hint,
            start_row_hint=start_row_hint,
        )
    except Exception as exc:
        raise _phase_error(f"Read workbook sheet structure ({_project_rel(workbook_path)})", exc) from exc
    try:
        target_paddle = collect_target_paddle_pair(release, form)
    except (FileNotFoundError, ValueError, XlsxMappingAutoGenerationError):
        raise
    except Exception as exc:
        raise _phase_error("Collect target template OCR context", exc) from exc
    try:
        target_roi = load_target_roi_pair(release, form)
    except (FileNotFoundError, ValueError):
        raise
    except Exception as exc:
        raise _phase_error("Load target ROI schemas", exc) from exc
    try:
        example = load_example_package()
    except (FileNotFoundError, ValueError):
        raise
    except Exception as exc:
        raise _phase_error("Load example XLSX mapping package", exc) from exc

    schema = _generation_schema()
    payload = {
        "target": {
            "release": release,
            "form": form,
            "template_workbook_path": _project_rel(workbook_path),
            "resolved_sheet": target_xlsx["sheet_name"],
            "resolved_start_row": target_xlsx["suggested_start_row"],
        },
        "target_paddle_ocr_from_both_template_images": target_paddle,
        "target_roi_schemas": target_roi,
        "target_xlsx_sheet_structure": target_xlsx,
        "example_package": example,
        "previous_validation_feedback": validation_feedback or [],
    }

    cfg = XLSX_MAPPING_AUTO_GENERATE if isinstance(XLSX_MAPPING_AUTO_GENERATE, dict) else {}
    use_openai = bool(cfg.get("use_openai", True))
    instructions = _instructions()
    if use_openai:
        model_override = str(cfg.get("openai_model") or "").strip() or None
        reasoning_enabled = bool(cfg.get("openai_reasoning", True))
        reasoning_cfg = None if reasoning_enabled else {"effort": "none"}
        provider = "openai"
        model_name = model_override or resolve_model(model_profile=str(cfg.get("openai_model_profile", "max_quality") or "max_quality"))
        try:
            raw_obj = call_json(
                json.dumps(payload, ensure_ascii=False),
                schema=schema,
                schema_name="xlsx_mapping_generation",
                instructions=instructions,
                model=model_override,
                model_profile=None if model_override else str(cfg.get("openai_model_profile", "max_quality") or "max_quality"),
                max_output_tokens=int(cfg.get("max_output_tokens", 24000) or 24000),
                reasoning=reasoning_cfg,
            )
        except Exception as exc:
            raise _phase_error(f"OpenAI structured mapping call (model={model_name})", exc) from exc
    else:
        provider = "local"
        model_name = str(cfg.get("local_model", "qwen3.5:9b") or "qwen3.5:9b")
        try:
            raw_obj = _call_local_json(payload, schema, instructions)
        except Exception as exc:
            raise _phase_error(f"Local structured mapping call (model={model_name})", exc) from exc

    try:
        mapping = _normalize_generated_mapping(
            raw_obj,
            form=form,
            fallback_sheet=str(target_xlsx["sheet_name"]),
            fallback_start_row=int(target_xlsx["suggested_start_row"]),
        )
        workbook_issues = _workbook_mapping_issues(mapping, target_xlsx)
    except Exception as exc:
        raise _phase_error("Normalize/validate generated mapping", exc) from exc
    return {
        "mapping": mapping,
        "raw_model_output": raw_obj,
        "workbook_issues": workbook_issues,
        "provider": provider,
        "model": model_name,
        "template_workbook_path": _project_rel(workbook_path),
        "target_sheet": target_xlsx["sheet_name"],
        "target_start_row": target_xlsx["suggested_start_row"],
        "target_paddle_line_count": sum(
            int((side.get("document_structure_from_paddleocr") or {}).get("line_count", 0) or 0)
            for side in target_paddle.get("sides", {}).values()
        ),
        "target_roi_source_count": len(target_roi.get("all_sources", [])),
        "example_mapping_path": example.get("mapping_path"),
        "confidence_notes": raw_obj.get("confidence_notes", []),
    }
