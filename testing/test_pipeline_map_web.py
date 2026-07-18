#!/usr/bin/env python3
"""Smoke tests for the ROI-to-XLSX visualizer graph builder and API."""

from __future__ import annotations

import base64
import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "py"))

import pipeline_map_web as pm


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def _write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2), encoding="utf-8")


def _build_fixture(root: Path) -> Path:
    release = "2099"
    template_name = "fixture template"
    form = "tpre"
    for side in ("a", "b"):
        tpl = root / "data" / "templates" / release / f"{form}_{side}.png"
        tpl.parent.mkdir(parents=True, exist_ok=True)
        tpl.write_bytes(PNG_1X1)

    schema_a = {
        "image_width": 100,
        "image_height": 100,
        "rois": [
            {"name": "id", "x": 2, "y": 2, "w": 20, "h": 8},
            {"name": "1", "x": 10, "y": 10, "w": 20, "h": 20},
            {"name": "1a", "x": 10, "y": 40, "w": 8, "h": 8},
            {"name": "1b", "x": 24, "y": 40, "w": 8, "h": 8},
            {"name": "2", "x": 50, "y": 10, "w": 20, "h": 20},
        ],
    }
    schema_b = {
        "image_width": 100,
        "image_height": 100,
        "rois": [
            {"name": "3", "x": 10, "y": 10, "w": 20, "h": 20},
        ],
    }
    _write_json(root / "data" / "roi_schemas" / release / f"{form}_a.json", schema_a)
    _write_json(root / "data" / "roi_schemas" / release / f"{form}_b.json", schema_b)
    mapping = {
        "form_type": form,
        "sheet": "Data",
        "start_row": 4,
        "mappings": [
            {"source": "id", "column": "A", "type": "direct"},
            {"source": "1", "column": "B", "type": "lookup", "map": {"a": "No", "b": "Yes"}},
            {"source": "2", "type": "multi_column", "mark": "Yes", "choices": {"a": "C", "b": "D"}},
            {"column": "E", "type": "static", "value": "NPS"},
            {"source": "missing", "column": "Z", "type": "direct"},
            {"source": "page_odd", "column": "F", "type": "direct"},
        ],
    }
    _write_json(root / "data" / "xlsx" / "mappings" / template_name / f"{form}.json", mapping)

    try:
        from openpyxl import Workbook
    except Exception as exc:  # pragma: no cover
        raise RuntimeError(f"openpyxl required for this test: {exc}") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for idx, value in enumerate(["ID", "Lookup", "Choice A", "Choice B", "Static", "Odd Page"], start=1):
        ws.cell(row=1, column=idx, value=value)
    ws.merge_cells("C1:D1")
    ws["C1"] = "Choices"
    workbook = root / "fixture.xlsx"
    wb.save(workbook)
    return workbook


def _patch_roots(root: Path) -> dict[str, object]:
    old = {
        "PROJECT_ROOT": pm.PROJECT_ROOT,
        "DATA_ROOT": pm.DATA_ROOT,
        "TEMPLATES_ROOT": pm.TEMPLATES_ROOT,
        "ROI_SCHEMAS_ROOT": pm.ROI_SCHEMAS_ROOT,
        "MAPPINGS_ROOT": pm.MAPPINGS_ROOT,
        "DATA_RELEASE": pm.DATA_RELEASE,
        "ACTIVE_XLSX_TEMPLATE_NAME": pm.ACTIVE_XLSX_TEMPLATE_NAME,
    }
    pm.PROJECT_ROOT = root
    pm.DATA_ROOT = root / "data"
    pm.TEMPLATES_ROOT = root / "data" / "templates"
    pm.ROI_SCHEMAS_ROOT = root / "data" / "roi_schemas"
    pm.MAPPINGS_ROOT = root / "data" / "xlsx" / "mappings"
    pm.DATA_RELEASE = "2099"
    pm.ACTIVE_XLSX_TEMPLATE_NAME = "fixture template"
    return old


def _restore_roots(old: dict[str, object]) -> None:
    for key, value in old.items():
        setattr(pm, key, value)


def main() -> int:
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        workbook = _build_fixture(root)
        old = _patch_roots(root)
        try:
            graph = pm.build_visualization("2099", "tpre", workbook_path=workbook)
            assert graph["stats"]["mapping_count"] == 6
            assert graph["sheet_grid"]["workbook_backed"] is True
            assert graph["sheet_grid"]["start_row"] == 4
            assert any(
                cell["value"] == "Lookup"
                for col in graph["sheet_grid"]["columns"]
                if col["column"] == "B"
                for cell in col["cells"]
            )
            assert graph["sheet_grid"]["visible_rows"] == [1, 2, 4]
            assert any(r["range"] == "C1:D1" and r["colspan"] == 2 for r in graph["sheet_grid"]["visible_merged_ranges"])
            assert any(
                cell["value"] == "Choices" and cell["merged"].get("colspan") == 2
                for col in graph["sheet_grid"]["columns"]
                if col["column"] == "C"
                for cell in col["cells"]
            )
            assert any(n["kind"] == "mcq" and n["name"] == "1" for n in graph["roi_nodes"])
            assert any(n["kind"] == "mcq_choice" and n["name"] == "1a" for n in graph["roi_nodes"])
            assert any(n["id"] == "virtual:id" and n["mapped"] for n in graph["virtual_nodes"])
            assert any(n["id"] == "virtual:page_odd" and n["mapped"] for n in graph["virtual_nodes"])
            assert any(n["id"] == "static:4" for n in graph["static_nodes"])
            assert any(e["mapping_type"] == "multi_column" and e["to"] == "col:C" for e in graph["edges"])
            assert any(e["mapping_type"] == "direct" and e["from"] == "virtual:id" and e["to"] == "cell:A:4" for e in graph["direct_edges"])
            assert any(e["mapping_type"] == "direct" and e["from"] == "virtual:page_odd" and e["to"] == "cell:F:4" for e in graph["direct_edges"])
            assert any(e["mapping_type"] == "lookup" and e["from"] == "roi:a:1" and e["to"] == "cell:B:4" for e in graph["direct_edges"])
            assert sum(1 for e in graph["direct_edges"] if e["mapping_type"] == "multi_column" and e["from"] == "roi:a:2") == 2
            assert any(e["mapping_type"] == "static" and e["from"] == "static:4" and e["to"] == "cell:E:4" for e in graph["direct_edges"])
            assert any(i["code"] == "missing_source" for i in graph["issues"])
            assert any(i["code"] == "missing_column" for i in graph["issues"])

            fallback_graph = pm.build_visualization("2099", "tpre", workbook_path=None)
            assert fallback_graph["sheet_grid"]["workbook_backed"] is False
            assert any(
                cell["id"] == "cell:A:4"
                for col in fallback_graph["sheet_grid"]["columns"]
                for cell in col["cells"]
            )

            app = pm.create_app()
            client = app.test_client()
            resp = client.get("/api/options?release=2099&form=tpre")
            assert resp.status_code == 200
            assert resp.get_json()["selected_release"] == "2099"
            resp = client.get(f"/api/visualization?release=2099&form=tpre&workbook={workbook}")
            assert resp.status_code == 200
            assert resp.get_json()["stats"]["mapping_count"] == 6
            assert resp.get_json()["direct_edges"]
            assert resp.get_json()["sheet_grid"]["columns"]
            resp = client.get("/api/template-image?release=2099&form=tpre&side=a")
            assert resp.status_code == 200
        finally:
            _restore_roots(old)
    print("pipeline_map_web tests passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
