#!/usr/bin/env python3
"""Tests for explicit text typing in XLSX mapped cells."""

from __future__ import annotations

import sys
import unittest
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "py"))

from xlsx_data_entry import fill_row  # noqa: E402


class XlsxTextCellTests(unittest.TestCase):
    def test_force_text_cells_converts_values_and_sets_excel_text_type(self) -> None:
        wb = Workbook()
        ws = wb.active
        mapping = {
            "mappings": [
                {"type": "static", "column": "A", "value": 123},
                {"type": "static", "column": "B", "value": "=1+1"},
                {"type": "direct", "column": "C", "source": "score", "transform": "number"},
            ]
        }
        item = {"data": [{"name": "score", "text": "45.5"}]}

        fill_row(ws, 2, mapping, item, cfg={"force_text_cells": True})

        self.assertEqual([ws["A2"].value, ws["B2"].value, ws["C2"].value], ["123", "=1+1", "45.5"])
        for cell in (ws["A2"], ws["B2"], ws["C2"]):
            self.assertEqual(cell.data_type, "s")
            self.assertEqual(cell.number_format, "@")
        wb.close()

    def test_force_text_cells_can_be_disabled(self) -> None:
        wb = Workbook()
        ws = wb.active
        mapping = {"mappings": [{"type": "static", "column": "A", "value": 123}]}

        fill_row(ws, 2, mapping, {}, cfg={"force_text_cells": False})

        self.assertEqual(ws["A2"].value, 123)
        self.assertEqual(ws["A2"].data_type, "n")
        self.assertEqual(ws["A2"].number_format, "General")
        wb.close()

    def test_blank_text_cell_gets_comment_with_fusion_raw_sources(self) -> None:
        wb = Workbook()
        ws = wb.active
        mapping = {
            "mappings": [
                {
                    "type": "direct",
                    "column": "A",
                    "source": "25",
                    "comment": {"text": "Raw: {raw}\nOCR confidence: {ocr_confidence}"},
                }
            ]
        }
        item = {
            "data": [
                {
                    "name": "25",
                    "kind": "text",
                    "text": "",
                    "_raw_ocr_text": "VLM optional answer",
                    "ocr_workflow": "paddle_vlm_fusion",
                    "ocr_fusion_text": "VLM optional answer Paddle optional answer",
                    "ocr_confidence_label": "medium",
                    "ocr_confidence_score": 0.81234,
                    "ocr_paddle_confidence": {
                        "detected_text": "Paddle optional answer",
                    },
                }
            ]
        }

        fill_row(ws, 2, mapping, item, cfg={"force_text_cells": True})

        self.assertEqual(ws["A2"].value, "")
        self.assertIsNotNone(ws["A2"].comment)
        self.assertEqual(
            ws["A2"].comment.text,
            "Raw: 'Paddle optional answer' + 'VLM optional answer' = 'VLM optional answer Paddle optional answer'\nOCR confidence: medium (0.812)",
        )
        wb.close()

    def test_text_comment_raw_uses_individual_ocr_source_without_fusion(self) -> None:
        wb = Workbook()
        ws = wb.active
        mapping = {
            "mappings": [
                {
                    "type": "direct",
                    "column": "A",
                    "source": "13",
                    "comment": {"text": "Raw: {raw}"},
                },
                {
                    "type": "direct",
                    "column": "B",
                    "source": "14",
                    "comment": {"text": "Raw: {raw}"},
                },
            ]
        }
        item = {
            "data": [
                {
                    "name": "13",
                    "kind": "text",
                    "text": "F",
                    "_raw_ocr_text": "vision urethra",
                    "ocr_text_source": "vision",
                    "ocr_paddle_confidence": {"detected_text": "paddle urethra"},
                },
                {
                    "name": "14",
                    "kind": "text",
                    "text": "B",
                    "_raw_ocr_text": "vision labia",
                    "ocr_text_source": "paddle",
                    "ocr_paddle_confidence": {"detected_text": "paddle labia"},
                },
            ]
        }

        fill_row(ws, 2, mapping, item, cfg={"force_text_cells": True})

        self.assertEqual(ws["A2"].comment.text, "Raw: vision urethra")
        self.assertEqual(ws["B2"].comment.text, "Raw: paddle labia")
        wb.close()


if __name__ == "__main__":
    unittest.main()
