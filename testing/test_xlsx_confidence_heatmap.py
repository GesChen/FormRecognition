#!/usr/bin/env python3
"""Tests for XLSX text ROI confidence heatmap fills."""

from __future__ import annotations

import sys
import unittest
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "py"))

from xlsx_data_entry import fill_row  # noqa: E402


def _rgb(cell) -> str:
    return str(cell.fill.fgColor.rgb or "")[-6:]


class XlsxConfidenceHeatmapTests(unittest.TestCase):
    def test_text_roi_confidence_controls_fill(self) -> None:
        wb = Workbook()
        ws = wb.active
        mapping = {
            "mappings": [
                {"source": "10", "column": "A", "type": "direct"},
                {"source": "11", "column": "B", "type": "direct"},
            ],
        }
        item = {
            "data": [
                {"name": "10", "kind": "text", "text": "low", "ocr_confidence_score": 0.0},
                {"name": "11", "kind": "text", "text": "high", "ocr_confidence_score": 1.0},
            ],
        }
        cfg = {"confidence_heatmap": {"enabled": True, "max_red": "F4B6B6"}}

        fill_row(ws, 2, mapping, item, cfg=cfg)

        self.assertEqual(_rgb(ws["A2"]), "F4B6B6")
        self.assertEqual(_rgb(ws["B2"]), "FFFFFF")
        self.assertEqual(ws["A2"].fill.fgColor.rgb, "FFF4B6B6")
        self.assertEqual(ws["B2"].fill.fgColor.rgb, "FFFFFFFF")
        wb.close()

    def test_heatmap_ignores_disabled_and_non_text_rois(self) -> None:
        wb = Workbook()
        ws = wb.active
        mapping = {
            "mappings": [
                {"source": "1", "column": "A", "type": "direct"},
                {"source": "2", "column": "B", "type": "direct"},
            ],
        }
        item = {
            "data": [
                {"name": "1", "kind": "text", "text": "off", "ocr_confidence_score": 0.0},
                {"name": "2", "kind": "mcq", "text": "a", "ocr_confidence_score": 0.0},
            ],
        }

        fill_row(ws, 2, mapping, item, cfg={"confidence_heatmap": {"enabled": False}})
        fill_row(ws, 3, mapping, item, cfg={"confidence_heatmap": {"enabled": True}})

        self.assertIsNone(ws["A2"].fill.fill_type)
        self.assertIsNone(ws["B3"].fill.fill_type)
        wb.close()


if __name__ == "__main__":
    unittest.main()
