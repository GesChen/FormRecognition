from __future__ import annotations

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "py"))

from openpyxl import Workbook, load_workbook

import form_sheet_map
from xlsx_data_entry import fill_from_pipeline, original_sheet_name


def _fixture(tmp_path: Path) -> dict:
    release = "2099"
    data_root = tmp_path / "data"
    mapping_dir = data_root / "xlsx" / "mappings" / release
    mapping_dir.mkdir(parents=True)
    mapping = {
        "form_type": "6post",
        "sheet": "6th Grade Post-Assessment Data",
        "start_row": 2,
        "mappings": [
            {"type": "direct", "source": "id", "column": "A"},
            {"type": "direct", "source": "teacher", "column": "B"},
        ],
    }
    (mapping_dir / "6post.json").write_text(json.dumps(mapping), encoding="utf-8")

    template_dir = data_root / "xlsx" / "workbook_templates" / release
    template_dir.mkdir(parents=True)
    template = template_dir / "2099 template.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = mapping["sheet"]
    ws["A1"] = "ID"
    ws["B1"] = "Teacher"
    ws["A2"] = "template"
    ws["B2"] = "template"
    ws.row_dimensions[2].height = 33
    wb.save(template)
    form_sheet_map.PATHS = {"data": data_root}

    return {
        "mapping_dir": mapping_dir,
        "output_dir": tmp_path,
        "staging_dir": tmp_path / "staging",
    }


def _items(teacher: str) -> list[dict]:
    return [
        {
            "id": "123",
            "form_type": "6post",
            "data": [{"name": "teacher", "kind": "text", "text": teacher}],
        }
    ]


def test_fill_from_pipeline_keeps_single_sheet_without_original_items(tmp_path):
    cfg = _fixture(tmp_path)
    out = fill_from_pipeline(_items("Peterson"), pdf_stem="single", cfg=cfg, verbose=False)

    wb = load_workbook(out)
    try:
        assert wb.sheetnames == ["6th Grade Post-Assessment Data"]
        ws = wb["6th Grade Post-Assessment Data"]
        assert ws["A2"].value == "123"
        assert ws["B2"].value == "Peterson"
        assert ws.freeze_panes == "A2"
    finally:
        wb.close()


def test_fill_from_pipeline_writes_original_sheet_pair(tmp_path):
    cfg = _fixture(tmp_path)
    out = fill_from_pipeline(
        _items("Peterson"),
        pdf_stem="paired",
        cfg=cfg,
        verbose=False,
        original_items=_items("Ms Petterson"),
    )

    wb = load_workbook(out)
    try:
        original_name = original_sheet_name("6post")
        assert wb.sheetnames == ["6th Grade Post-Assessment Data", original_name]
        assert "(" in original_name and ")" in original_name
        assert len(original_name) <= 31
        assert wb["6th Grade Post-Assessment Data"]["B2"].value == "Peterson"
        assert wb["6th Grade Post-Assessment Data"].freeze_panes == "A2"
        assert wb[original_name]["B2"].value == "Ms Petterson"
        assert wb[original_name].row_dimensions[2].height == 33
        assert wb[original_name].freeze_panes == "A2"
    finally:
        wb.close()
