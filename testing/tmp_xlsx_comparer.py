#!/usr/bin/env python3
"""
Temporary XLSX comparer for truth-vs-test sheet evaluation.

Compares two worksheets row-by-row (same-order assumption), and reports:
  - compares only rows present in the test workbook (non-empty test ID rows)
  - row coverage/alignment
  - ID recognition metrics (precision/recall/F1, row-accuracy, duplicates, format validity)
  - per-cell and per-column exact-match metrics
  - completeness metrics (miss/overfill)
  - categorical/MCQ-like accuracy
  - text-like fuzzy similarity

Usage example:
  python3 testing/tmp_xlsx_comparer.py \
    --truth-file data/truth.xlsx --truth-sheet "6th Grade Pre-Assessment Data" \
    --test-file output/xlsx/run.xlsx --test-sheet "6th Grade Pre-Assessment Data"
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string, get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "openpyxl is required. Install with: pip install openpyxl"
    ) from exc


LIKERT_VALUES = {
    "strongly disagree",
    "disagree",
    "neither agree nor disagree",
    "agree",
    "strongly agree",
}
GRADE_VALUES = {
    "6th grade",
    "7th grade",
    "8th grade",
    "9th grade",
    "other grade",
}
GENDER_VALUES = {
    "girl",
    "boy",
    "non-binary",
    "unknown/not reported",
}


@dataclass
class CompareRange:
    start_row: int
    end_row: int
    start_col: int
    end_col: int


def _normalize_cell(value: Any) -> str:
    """Normalize cell content for stable comparisons."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return f"{value:.12g}"
    s = str(value).strip()
    s = re.sub(r"\s+", " ", s)
    return s


def _safe_div(n: float, d: float) -> float | None:
    return (n / d) if d else None


def _f1(precision: float | None, recall: float | None) -> float | None:
    if precision is None or recall is None:
        return None
    if precision + recall == 0:
        return 0.0
    return 2.0 * precision * recall / (precision + recall)


def _levenshtein_distance(a: str, b: str) -> int:
    """Classic DP Levenshtein distance, memory-optimized."""
    if a == b:
        return 0
    if not a:
        return len(b)
    if not b:
        return len(a)
    if len(a) > len(b):
        a, b = b, a
    prev = list(range(len(a) + 1))
    for i, cb in enumerate(b, start=1):
        cur = [i]
        for j, ca in enumerate(a, start=1):
            ins = cur[j - 1] + 1
            delete = prev[j] + 1
            sub = prev[j - 1] + (0 if ca == cb else 1)
            cur.append(min(ins, delete, sub))
        prev = cur
    return prev[-1]


def _similarity(a: str, b: str) -> float:
    denom = max(len(a), len(b), 1)
    return 1.0 - (_levenshtein_distance(a, b) / denom)


def _is_row_nonempty(row_vals: list[str]) -> bool:
    return any(v != "" for v in row_vals)


def _last_nonempty_row(ws, start_row: int, start_col: int, end_col: int) -> int:
    last = start_row
    max_row = int(ws.max_row or start_row)
    for r in range(start_row, max_row + 1):
        row_vals = [_normalize_cell(ws.cell(r, c).value) for c in range(start_col, end_col + 1)]
        if _is_row_nonempty(row_vals):
            last = r
    return last


def _last_nonempty_id_row(ws, start_row: int, id_col_idx: int) -> int:
    """Last row with a non-empty ID value in id_col_idx."""
    last = start_row
    max_row = int(ws.max_row or start_row)
    for r in range(start_row, max_row + 1):
        if _normalize_cell(ws.cell(r, id_col_idx).value):
            last = r
    return last


def _classify_categorical(v: str) -> tuple[str, str] | None:
    if not v:
        return None
    low = v.strip().lower()
    if re.fullmatch(r"[a-h]", low):
        return ("letter", low)
    if low in {"true", "false"}:
        return ("boolean", low)
    if low in LIKERT_VALUES:
        return ("likert", low)
    if low in GRADE_VALUES:
        return ("grade", low)
    if low in GENDER_VALUES:
        return ("gender", low)
    return None


def _is_text_like(v: str) -> bool:
    if not v:
        return False
    if _classify_categorical(v):
        return False
    return bool(re.search(r"[A-Za-z]", v))


def _serializable_none_to_null(d: dict[str, Any]) -> dict[str, Any]:
    return json.loads(json.dumps(d, default=str))


def _build_truth_id_index(
    wb_truth,
    *,
    start_row: int,
    id_col_idx: int,
) -> tuple[dict[str, list[tuple[str, int]]], int]:
    """
    Build lookup: normalized ID -> [(sheet_name, row_index), ...] across all truth sheets.
    Returns (index, total_nonempty_id_rows).
    """
    idx: dict[str, list[tuple[str, int]]] = defaultdict(list)
    total = 0
    for sheet in wb_truth.sheetnames:
        ws = wb_truth[sheet]
        max_row = int(ws.max_row or start_row)
        for r in range(start_row, max_row + 1):
            tid = _normalize_cell(ws.cell(r, id_col_idx).value).upper()
            if not tid:
                continue
            idx[tid].append((sheet, r))
            total += 1
    return idx, total


def compare(
    *,
    truth_file: Path,
    truth_sheet: str | None,
    test_file: Path,
    test_sheet: str,
    compare_range: CompareRange,
    id_col_idx: int,
    id_regex: re.Pattern[str],
    id_match_any_truth_sheet: bool = False,
) -> dict[str, Any]:
    wb_truth = load_workbook(truth_file, data_only=True)
    wb_test = load_workbook(test_file, data_only=True)

    if not id_match_any_truth_sheet and truth_sheet not in wb_truth.sheetnames:
        raise ValueError(f"truth sheet not found: {truth_sheet!r}")
    if test_sheet not in wb_test.sheetnames:
        raise ValueError(f"test sheet not found: {test_sheet!r}")

    ws_truth = wb_truth[truth_sheet] if (truth_sheet and not id_match_any_truth_sheet) else None
    ws_test = wb_test[test_sheet]

    start_row = compare_range.start_row
    end_row = compare_range.end_row
    start_col = compare_range.start_col
    end_col = compare_range.end_col

    per_col = {
        get_column_letter(c): {
            "compared_cells": 0,
            "exact_matches": 0,
            "truth_nonempty": 0,
            "test_nonempty": 0,
            "mismatch_examples": [],
        }
        for c in range(start_col, end_col + 1)
    }

    row_truth_nonempty = 0
    row_test_nonempty = 0
    row_both_nonempty = 0
    row_presence_match = 0
    rows_scanned_total = end_row - start_row + 1
    rows_compared = 0

    cell_compared = 0
    cell_exact = 0
    miss_count = 0  # truth has value, test is empty
    overfill_count = 0  # truth empty, test has value

    categorical_total = 0
    categorical_exact = 0
    categorical_by_type = defaultdict(lambda: {"total": 0, "exact": 0})
    letter_confusions = Counter()

    text_total = 0
    text_exact = 0
    text_similarity_sum = 0.0

    truth_ids: list[str] = []
    test_ids: list[str] = []
    # In id_match_any_truth_sheet mode we still track all non-empty test IDs
    # (matched + unmatched) so reporting denominators remain meaningful.
    test_ids_all: list[str] = []
    id_row_total = 0
    id_row_exact = 0
    id_truth_valid = 0
    id_test_valid = 0
    truth_lookup, truth_lookup_rows_total = ({}, 0)
    truth_lookup_found_count = 0
    truth_lookup_not_found_count = 0
    truth_lookup_unique_found: set[str] = set()
    truth_lookup_multi_match_ids: set[str] = set()
    if id_match_any_truth_sheet:
        truth_lookup, truth_lookup_rows_total = _build_truth_id_index(
            wb_truth,
            start_row=start_row,
            id_col_idx=id_col_idx,
        )
        for t, locs in truth_lookup.items():
            truth_ids.extend([t] * len(locs))
            id_truth_valid += int(bool(id_regex.fullmatch(t))) * len(locs)
        for t, locs in truth_lookup.items():
            if len(locs) > 1:
                truth_lookup_multi_match_ids.add(t)

    for r in range(start_row, end_row + 1):
        test_row_vals = [
            _normalize_cell(ws_test.cell(r, c).value) for c in range(start_col, end_col + 1)
        ]
        if id_match_any_truth_sheet:
            truth_row_vals = [""] * (end_col - start_col + 1)
        else:
            truth_row_vals = [
                _normalize_cell(ws_truth.cell(r, c).value) for c in range(start_col, end_col + 1)
            ]

        truth_nonempty = _is_row_nonempty(truth_row_vals)
        test_nonempty = _is_row_nonempty(test_row_vals)

        row_truth_nonempty += int(truth_nonempty)
        row_test_nonempty += int(test_nonempty)
        row_both_nonempty += int(truth_nonempty and test_nonempty)
        row_presence_match += int(truth_nonempty == test_nonempty)

        pid = _normalize_cell(ws_test.cell(r, id_col_idx).value).upper()
        if id_match_any_truth_sheet:
            tid = ""
            if pid:
                hits = truth_lookup.get(pid, [])
                if hits:
                    truth_lookup_found_count += 1
                    truth_lookup_unique_found.add(pid)
                    if len(hits) > 1:
                        truth_lookup_multi_match_ids.add(pid)
                    tid_sheet, tid_row = hits[0]
                    ws_truth_hit = wb_truth[tid_sheet]
                    tid = _normalize_cell(ws_truth_hit.cell(tid_row, id_col_idx).value).upper()
                    truth_row_vals = [
                        _normalize_cell(ws_truth_hit.cell(tid_row, c).value)
                        for c in range(start_col, end_col + 1)
                    ]
                else:
                    truth_lookup_not_found_count += 1
                    truth_row_vals = [""] * (end_col - start_col + 1)
                    tid = ""
            truth_nonempty = _is_row_nonempty(truth_row_vals)
        else:
            tid = _normalize_cell(ws_truth.cell(r, id_col_idx).value).upper()

        # Track total non-empty test IDs before any matching filters.
        if pid:
            test_ids_all.append(pid)

        # Compare only rows present in test workbook (non-empty test ID).
        if not pid:
            continue
        # In ID-match mode, exclude unmatched IDs from comparison metrics.
        if id_match_any_truth_sheet and not tid:
            continue
        rows_compared += 1

        if tid and not id_match_any_truth_sheet:
            truth_ids.append(tid)
            id_truth_valid += int(bool(id_regex.fullmatch(tid)))
        if pid:
            test_ids.append(pid)
            id_test_valid += int(bool(id_regex.fullmatch(pid)))

        if tid or pid:
            id_row_total += 1
            id_row_exact += int(tid == pid)

        for c in range(start_col, end_col + 1):
            col = get_column_letter(c)
            idx = c - start_col
            t = truth_row_vals[idx] if 0 <= idx < len(truth_row_vals) else ""
            p = test_row_vals[idx] if 0 <= idx < len(test_row_vals) else ""

            pc = per_col[col]
            if t:
                pc["truth_nonempty"] += 1
            if p:
                pc["test_nonempty"] += 1

            if not (t or p):
                continue

            cell_compared += 1
            pc["compared_cells"] += 1
            exact = (t == p)
            cell_exact += int(exact)
            pc["exact_matches"] += int(exact)

            if not exact and len(pc["mismatch_examples"]) < 3:
                pc["mismatch_examples"].append({"row": r, "truth": t, "test": p})

            if t and not p:
                miss_count += 1
            elif p and not t:
                overfill_count += 1

            t_cat = _classify_categorical(t)
            p_cat = _classify_categorical(p)
            if t_cat or p_cat:
                # Compare on union of "categorical-like" cells.
                categorical_total += 1
                categorical_exact += int(t.lower() == p.lower())
                cat_type = (t_cat or p_cat)[0]
                categorical_by_type[cat_type]["total"] += 1
                categorical_by_type[cat_type]["exact"] += int(t.lower() == p.lower())
                if cat_type == "letter" and t and p and t.lower() != p.lower():
                    letter_confusions[f"{t.lower()}->{p.lower()}"] += 1

            if _is_text_like(t) or _is_text_like(p):
                text_total += 1
                text_exact += int(t.lower() == p.lower())
                text_similarity_sum += _similarity(t.lower(), p.lower())

    truth_id_set = set(truth_ids)
    test_id_set = set(test_ids)
    test_id_set_all = set(test_ids_all)
    id_intersection = truth_id_set & test_id_set

    if id_match_any_truth_sheet:
        id_precision = _safe_div(truth_lookup_found_count, len(test_ids_all))
        id_recall = _safe_div(len(truth_lookup_unique_found), len(truth_id_set))
    else:
        id_precision = _safe_div(len(id_intersection), len(test_id_set))
        id_recall = _safe_div(len(id_intersection), len(truth_id_set))
    id_f1 = _f1(id_precision, id_recall)

    truth_dup_count = sum(cnt - 1 for cnt in Counter(truth_ids).values() if cnt > 1)
    test_dup_count = sum(cnt - 1 for cnt in Counter(test_ids).values() if cnt > 1)

    # Finish per-column accuracy fields
    for stats in per_col.values():
        stats["accuracy"] = _safe_div(stats["exact_matches"], stats["compared_cells"])

    macro_values = [v["accuracy"] for v in per_col.values() if v["accuracy"] is not None]
    macro_acc = (sum(macro_values) / len(macro_values)) if macro_values else None

    top_confusions = [{"pair": k, "count": v} for k, v in letter_confusions.most_common(10)]

    rows_block = {
        "rows_scanned_total": rows_scanned_total,
        "total_rows_compared": rows_compared,
        "truth_nonempty_rows": row_truth_nonempty,
        "test_nonempty_rows": row_test_nonempty,
        "both_nonempty_rows": row_both_nonempty,
        "row_presence_match_rate": _safe_div(row_presence_match, rows_scanned_total),
    }
    if id_match_any_truth_sheet:
        rows_block.update(
            {
                "truth_nonempty_rows": None,
                "both_nonempty_rows": None,
                "row_presence_match_rate": None,
                "matched_by_id_rows": truth_lookup_found_count,
                "unmatched_by_id_rows": truth_lookup_not_found_count,
            }
        )

    result = {
        "inputs": {
            "truth_file": str(truth_file),
            "truth_sheet": truth_sheet,
            "test_file": str(test_file),
            "test_sheet": test_sheet,
            "start_row": start_row,
            "end_row": end_row,
            "start_col": get_column_letter(start_col),
            "end_col": get_column_letter(end_col),
            "id_column": get_column_letter(id_col_idx),
            "id_regex": id_regex.pattern,
            "row_filter": "test_id_nonempty_only",
            "id_match_any_truth_sheet": bool(id_match_any_truth_sheet),
            "truth_scope": "all_sheets_by_id" if id_match_any_truth_sheet else "single_sheet_row_aligned",
        },
        "rows": rows_block,
        "id_metrics": {
            "truth_nonempty_id_rows": truth_lookup_rows_total if id_match_any_truth_sheet else len(truth_ids),
            "test_nonempty_id_rows": (len(test_ids_all) if id_match_any_truth_sheet else len(test_ids)),
            "test_compared_id_rows": len(test_ids),
            "truth_unique_ids": len(truth_id_set),
            "test_unique_ids": (len(test_id_set_all) if id_match_any_truth_sheet else len(test_id_set)),
            "id_set_intersection": len(id_intersection),
            "id_precision": id_precision,
            "id_recall": id_recall,
            "id_f1": id_f1,
            "id_row_accuracy": _safe_div(id_row_exact, id_row_total),
            "id_row_total": id_row_total,
            "id_row_exact": id_row_exact,
            "test_ids_found_in_truth_count": truth_lookup_found_count if id_match_any_truth_sheet else len(id_intersection),
            "test_ids_found_in_truth_rate": _safe_div(
                truth_lookup_found_count if id_match_any_truth_sheet else len(id_intersection),
                (len(test_ids_all) if id_match_any_truth_sheet else len(test_ids)),
            ),
            "test_ids_not_found_in_truth_count": truth_lookup_not_found_count if id_match_any_truth_sheet else None,
            "truth_ids_with_multiple_rows_count": (
                len([k for k, v in (truth_lookup or {}).items() if len(v) > 1])
                if id_match_any_truth_sheet
                else None
            ),
            "truth_id_format_valid_rate": _safe_div(id_truth_valid, len(truth_ids)),
            "test_id_format_valid_rate": _safe_div(id_test_valid, len(test_ids)),
            "truth_duplicate_id_count": truth_dup_count,
            "test_duplicate_id_count": test_dup_count,
            "truth_duplicate_id_rate": _safe_div(truth_dup_count, len(truth_ids)),
            "test_duplicate_id_rate": _safe_div(test_dup_count, len(test_ids)),
        },
        "cell_metrics": {
            "cells_compared_nonempty_union": cell_compared,
            "cells_exact_match": cell_exact,
            "cell_micro_accuracy": _safe_div(cell_exact, cell_compared),
            "cell_macro_accuracy_by_column": macro_acc,
            "miss_count_truth_nonempty_test_empty": miss_count,
            "overfill_count_truth_empty_test_nonempty": overfill_count,
            "miss_rate": _safe_div(miss_count, cell_compared),
            "overfill_rate": _safe_div(overfill_count, cell_compared),
        },
        "categorical_metrics": {
            "overall_mcq_accuracy": _safe_div(
                categorical_by_type["letter"]["exact"],
                categorical_by_type["letter"]["total"],
            ),
            "categorical_cells_compared": categorical_total,
            "categorical_exact_match_accuracy": _safe_div(categorical_exact, categorical_total),
            "text_like_cells_compared": text_total,
            "text_like_exact_match_accuracy_case_insensitive": _safe_div(text_exact, text_total),
            "text_like_avg_similarity_case_insensitive": _safe_div(text_similarity_sum, text_total),
            "by_type": {
                k: {
                    "total": v["total"],
                    "exact": v["exact"],
                    "accuracy": _safe_div(v["exact"], v["total"]),
                }
                for k, v in sorted(categorical_by_type.items())
            },
            "letter_confusions_top10": top_confusions,
        },
        "mcq_metrics": {
            "mcq_cells_compared": categorical_by_type["letter"]["total"],
            "mcq_cells_exact": categorical_by_type["letter"]["exact"],
            "overall_mcq_accuracy": _safe_div(
                categorical_by_type["letter"]["exact"],
                categorical_by_type["letter"]["total"],
            ),
        },
        "text_metrics": {
            "text_like_cells_compared": text_total,
            "text_like_exact_match_accuracy_case_insensitive": _safe_div(text_exact, text_total),
            "text_like_avg_similarity_case_insensitive": _safe_div(text_similarity_sum, text_total),
        },
        "per_column": per_col,
    }
    return _serializable_none_to_null(result)


def _print_summary(result: dict[str, Any], *, show_column_details: bool) -> None:
    rows = result["rows"]
    idm = result["id_metrics"]
    cells = result["cell_metrics"]
    cats = result["categorical_metrics"]
    mcq = result.get("mcq_metrics", {})
    textm = result["text_metrics"]

    print("XLSX COMPARISON SUMMARY")
    truth_scope = result["inputs"].get("truth_scope")
    truth_sheet_disp = (
        "ALL_SHEETS_BY_ID" if truth_scope == "all_sheets_by_id" else result["inputs"].get("truth_sheet")
    )
    print(f"- Truth: {result['inputs']['truth_file']} [{truth_sheet_disp}]")
    print(f"- Test:  {result['inputs']['test_file']} [{result['inputs']['test_sheet']}]")
    print(
        f"- Range: rows {result['inputs']['start_row']}..{result['inputs']['end_row']}, "
        f"cols {result['inputs']['start_col']}..{result['inputs']['end_col']}"
    )
    print("")

    print("Row Metrics")
    print(f"- Row filter: {result['inputs'].get('row_filter', 'all_rows')}")
    print(f"- Rows scanned (range): {rows.get('rows_scanned_total')}")
    print(f"- Total rows compared: {rows['total_rows_compared']}")
    if rows.get("matched_by_id_rows") is not None:
        print(f"- Matched by ID rows: {rows.get('matched_by_id_rows')}")
        print(f"- Unmatched by ID rows: {rows.get('unmatched_by_id_rows')}")
    else:
        print(f"- Truth non-empty rows: {rows['truth_nonempty_rows']}")
    print(f"- Test non-empty rows:  {rows['test_nonempty_rows']}")
    if rows.get("both_nonempty_rows") is not None:
        print(f"- Both non-empty rows:  {rows['both_nonempty_rows']}")
    if rows.get("row_presence_match_rate") is not None:
        print(f"- Row presence match rate: {rows['row_presence_match_rate']}")
    print("")

    print("ID Metrics")
    if result["inputs"].get("id_match_any_truth_sheet"):
        print(f"- Test IDs found in truth: {idm.get('test_ids_found_in_truth_count')} / {idm.get('test_nonempty_id_rows')} "
              f"({idm.get('test_ids_found_in_truth_rate')})")
    print(f"- ID precision: {idm['id_precision']}")
    print(f"- ID recall:    {idm['id_recall']}")
    print(f"- ID F1:        {idm['id_f1']}")
    print(f"- ID row accuracy (same row): {idm['id_row_accuracy']}")
    print(
        f"- Truth/Test unique IDs: {idm['truth_unique_ids']} / {idm['test_unique_ids']} "
        f"(intersection {idm['id_set_intersection']})"
    )
    print(
        f"- Duplicate IDs (truth/test): {idm['truth_duplicate_id_count']} / "
        f"{idm['test_duplicate_id_count']}"
    )
    print("")

    print("Cell Metrics")
    print(f"- Cells compared (non-empty union): {cells['cells_compared_nonempty_union']}")
    print(f"- Cell micro accuracy: {cells['cell_micro_accuracy']}")
    print(f"- Cell macro accuracy (by column): {cells['cell_macro_accuracy_by_column']}")
    print(
        f"- Miss / Overfill counts: {cells['miss_count_truth_nonempty_test_empty']} / "
        f"{cells['overfill_count_truth_empty_test_nonempty']}"
    )
    print("")

    print("Categorical/Text Metrics")
    print(f"- MCQ overall accuracy: {cats.get('overall_mcq_accuracy')}")
    print(f"- Categorical accuracy: {cats['categorical_exact_match_accuracy']}")
    print(f"- Text-like exact (case-insensitive): {textm['text_like_exact_match_accuracy_case_insensitive']}")
    print(f"- Text-like avg similarity (case-insensitive): {textm['text_like_avg_similarity_case_insensitive']}")

    if show_column_details:
        print("")
        print("Per-Column Accuracy")
        for col, stats in result["per_column"].items():
            if stats["compared_cells"] <= 0:
                continue
            print(
                f"- {col}: accuracy={stats['accuracy']} "
                f"({stats['exact_matches']}/{stats['compared_cells']})"
            )


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Temporary XLSX comparer (truth vs test; compares only non-empty test ID rows). "
            "Supports row-order mode and ID-match-across-truth-sheets mode."
        ),
    )
    parser.add_argument("--truth-file", type=Path, required=True, help="Ground-truth workbook path.")
    parser.add_argument("--truth-sheet", required=False, default="", help="Sheet name in the truth workbook.")
    parser.add_argument("--test-file", type=Path, required=True, help="Predicted/test workbook path.")
    parser.add_argument("--test-sheet", required=True, help="Sheet name in the test workbook.")

    parser.add_argument(
        "--start-row",
        type=int,
        default=4,
        help="First row to compare (default: 4).",
    )
    parser.add_argument(
        "--end-row",
        type=int,
        default=None,
        help="Last row to compare (default: auto-detect last non-empty ID row in test sheet).",
    )
    parser.add_argument(
        "--start-col",
        default="A",
        help="First column to compare (default: A).",
    )
    parser.add_argument(
        "--end-col",
        default=None,
        help="Last column to compare (default: max used column across both sheets).",
    )
    parser.add_argument(
        "--id-column",
        default="A",
        help="Column containing ID values (default: A).",
    )
    parser.add_argument(
        "--id-regex",
        default=r"^ID:[A-Z0-9]{8}$",
        help="Regex used for ID format validity checks (default: ^ID:[A-Z0-9]{8}$).",
    )
    parser.add_argument(
        "--json-out",
        type=Path,
        default=None,
        help="Optional output path to write full JSON metrics.",
    )
    parser.add_argument(
        "--print-json",
        action="store_true",
        help="Print full JSON metrics to stdout.",
    )
    parser.add_argument(
        "--show-column-details",
        action="store_true",
        help="Print per-column accuracy lines in stdout summary.",
    )
    parser.add_argument(
        "--id-match-any-truth-sheet",
        action="store_true",
        help=(
            "Match each test-row ID against all truth workbook sheets. "
            "Disables row-aligned/sheet-specific truth comparison."
        ),
    )
    args = parser.parse_args()

    truth_file = args.truth_file.resolve()
    test_file = args.test_file.resolve()
    if not truth_file.is_file():
        print(f"Truth file not found: {truth_file}", file=sys.stderr)
        return 2
    if not test_file.is_file():
        print(f"Test file not found: {test_file}", file=sys.stderr)
        return 2

    try:
        start_col = column_index_from_string(str(args.start_col).strip().upper())
        id_col = column_index_from_string(str(args.id_column).strip().upper())
    except ValueError as exc:
        print(f"Invalid column argument: {exc}", file=sys.stderr)
        return 2

    wb_truth = load_workbook(truth_file, data_only=True)
    wb_test = load_workbook(test_file, data_only=True)
    if not args.id_match_any_truth_sheet and args.truth_sheet not in wb_truth.sheetnames:
        print(f"Truth sheet not found: {args.truth_sheet!r}", file=sys.stderr)
        return 2
    if args.test_sheet not in wb_test.sheetnames:
        print(f"Test sheet not found: {args.test_sheet!r}", file=sys.stderr)
        return 2
    ws_truth = wb_truth[args.truth_sheet] if not args.id_match_any_truth_sheet else None
    ws_test = wb_test[args.test_sheet]

    if args.end_col:
        try:
            end_col = column_index_from_string(str(args.end_col).strip().upper())
        except ValueError as exc:
            print(f"Invalid --end-col: {exc}", file=sys.stderr)
            return 2
    else:
        if ws_truth is None:
            end_col = int(ws_test.max_column or 1)
        else:
            end_col = max(int(ws_truth.max_column or 1), int(ws_test.max_column or 1))

    if end_col < start_col:
        print("--end-col must be >= --start-col", file=sys.stderr)
        return 2

    if args.end_row is not None:
        end_row = int(args.end_row)
    else:
        # Compare scope follows rows present in test workbook (non-empty ID rows).
        end_row = _last_nonempty_id_row(ws_test, args.start_row, id_col)
    if end_row < args.start_row:
        print("--end-row must be >= --start-row", file=sys.stderr)
        return 2

    try:
        id_regex = re.compile(str(args.id_regex), flags=re.IGNORECASE)
    except re.error as exc:
        print(f"Invalid --id-regex: {exc}", file=sys.stderr)
        return 2

    result = compare(
        truth_file=truth_file,
        truth_sheet=(args.truth_sheet or None),
        test_file=test_file,
        test_sheet=args.test_sheet,
        compare_range=CompareRange(
            start_row=int(args.start_row),
            end_row=end_row,
            start_col=start_col,
            end_col=end_col,
        ),
        id_col_idx=id_col,
        id_regex=id_regex,
        id_match_any_truth_sheet=bool(args.id_match_any_truth_sheet),
    )

    _print_summary(result, show_column_details=bool(args.show_column_details))

    if args.print_json:
        print("")
        print(json.dumps(result, indent=2, ensure_ascii=False))

    if args.json_out:
        out = args.json_out.resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        print(f"\nWrote JSON metrics: {out}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
